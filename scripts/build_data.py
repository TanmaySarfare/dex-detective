#!/usr/bin/env python3
"""Convert the Pokémon Excel workbook into data/pokemon.json for the game.

Usage:
    pip install -r scripts/requirements.txt
    python scripts/build_data.py data/source/Pokemon_Dataset_1.xlsx

The script reads columns by their header names, so column order in the
workbook doesn't matter. Generation and base stat total are worked out here
rather than read from the workbook's formula cells, so the result doesn't
depend on Excel having saved calculated values.
"""
import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

TYPES = {
    "Normal", "Fire", "Water", "Electric", "Grass", "Ice", "Fighting", "Poison", "Ground",
    "Flying", "Psychic", "Bug", "Rock", "Ghost", "Dragon", "Dark", "Steel", "Fairy",
}
# First National Pokédex number of each generation (Gen 1 to Gen 9).
GEN_STARTS = [1, 152, 252, 387, 494, 650, 722, 810, 906]

# JSON field -> spreadsheet header
COLUMNS = {
    "dex": "Dex No", "name": "Name", "category": "Category", "type1": "Type 1", "type2": "Type 2",
    "hp": "HP", "attack": "Attack", "defense": "Defense", "spAtk": "Sp. Atk", "spDef": "Sp. Def",
    "speed": "Speed", "heightM": "Height (m)", "weightKg": "Weight (kg)",
}
STATS = ["hp", "attack", "defense", "spAtk", "spDef", "speed"]
FIELDS = ["dex", "name", "generation", "category", "type1", "type2", *STATS, "total", "heightM", "weightKg"]


def generation(dex: int) -> int:
    return sum(1 for start in GEN_STARTS if dex >= start)


def parse_row(row, idx, line, problems):
    rec = {field: row[i] for field, i in idx.items()}
    try:
        rec["dex"] = int(rec["dex"])
        for s in STATS:
            rec[s] = int(rec[s])
        rec["heightM"] = round(float(rec["heightM"]), 1)
        rec["weightKg"] = round(float(rec["weightKg"]), 1)
    except (TypeError, ValueError):
        problems.append(f"row {line}: a number column is empty or not a number")
        return None
    rec["name"] = str(rec["name"] or "").strip()
    rec["category"] = str(rec["category"] or "").strip()
    rec["type1"] = str(rec["type1"] or "").strip()
    rec["type2"] = str(rec["type2"] or "").strip()

    if not rec["name"]:
        problems.append(f"row {line}: missing name")
    if rec["type1"] not in TYPES:
        problems.append(f"row {line} ({rec['name']}): unknown Type 1 '{rec['type1']}'")
    if rec["type2"] and (rec["type2"] not in TYPES or rec["type2"] == rec["type1"]):
        problems.append(f"row {line} ({rec['name']}): invalid Type 2 '{rec['type2']}'")
    if any(not 1 <= rec[s] <= 255 for s in STATS):
        problems.append(f"row {line} ({rec['name']}): a base stat is outside 1-255")
    if rec["heightM"] <= 0 or rec["weightKg"] <= 0:
        problems.append(f"row {line} ({rec['name']}): height and weight must be above 0")

    rec["generation"] = generation(rec["dex"])
    rec["total"] = sum(rec[s] for s in STATS)
    return [rec[f] for f in FIELDS]


def main() -> None:
    root = Path(__file__).resolve().parent.parent
    ap = argparse.ArgumentParser(description="Build data/pokemon.json from the Pokémon workbook.")
    ap.add_argument("workbook", type=Path, help="path to the .xlsx workbook")
    ap.add_argument("-o", "--output", type=Path, default=root / "data" / "pokemon.json")
    ap.add_argument("--sheet", help="sheet name (defaults to the first sheet)")
    args = ap.parse_args()

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]

    missing = [col for col in COLUMNS.values() if col not in header]
    if missing:
        sys.exit("Missing columns in the workbook: " + ", ".join(missing))
    idx = {field: header.index(col) for field, col in COLUMNS.items()}
    form_col = header.index("Form") if "Form" in header else None

    out, problems = [], []
    for line, row in enumerate(rows, start=2):
        if row[idx["dex"]] in (None, ""):
            continue
        # The game uses one entry per species, so alternate forms (Mega, regional...) are skipped.
        if form_col is not None and row[form_col] not in (None, "", "Base"):
            continue
        parsed = parse_row(row, idx, line, problems)
        if parsed:
            out.append(parsed)

    dex_numbers = [r[0] for r in out]
    if len(set(dex_numbers)) != len(dex_numbers):
        problems.append("the same Dex No appears more than once")
    if problems:
        print("Fix these problems in the workbook, then run the script again:", file=sys.stderr)
        for p in problems:
            print("  - " + p, file=sys.stderr)
        sys.exit(1)

    out.sort(key=lambda r: r[0])
    # One Pokémon per line keeps Git diffs readable when the data changes.
    body = ",\n".join(json.dumps(r, ensure_ascii=False, separators=(",", ":")) for r in out)
    text = '{"fields":' + json.dumps(FIELDS) + ',\n"rows":[\n' + body + "\n]}\n"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(text, encoding="utf-8")
    print(f"Wrote {len(out)} Pokémon to {args.output}")


if __name__ == "__main__":
    main()
